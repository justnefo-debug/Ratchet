import os
from pptx import Presentation
from typing import List, Dict
from core.models import Entity
from .base_handler import DocumentHandler

class PPTXHandler(DocumentHandler):
    def extract_text(self, file_path: str) -> str:
        prs = Presentation(file_path)
        text_parts = []
        
        for slide in prs.slides:
            # Extract notes
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                text_parts.append(slide.notes_slide.notes_text_frame.text)
                
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text_parts.append(shape.text)
                if shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            text_parts.append(cell.text)
                            
        return "\n".join(text_parts)
            
    def rebuild(self, file_path: str, mapping: Dict[str, str], entities: List[Entity]) -> str:
        prs = Presentation(file_path)
        
        def replace_in_text(text: str) -> str:
            new_text = text
            for placeholder, original in mapping.items():
                if original in new_text:
                    new_text = new_text.replace(original, placeholder)
            return new_text
            
        for slide in prs.slides:
            # Replace in notes
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_frame = slide.notes_slide.notes_text_frame
                for paragraph in notes_frame.paragraphs:
                    for run in paragraph.runs:
                        run.text = replace_in_text(run.text)
                        
            for shape in slide.shapes:
                if hasattr(shape, "text_frame") and shape.text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            run.text = replace_in_text(run.text)
                            
                if shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            for paragraph in cell.text_frame.paragraphs:
                                for run in paragraph.runs:
                                    run.text = replace_in_text(run.text)
                                    
        safe_path = file_path.replace(".pptx", "_safe.pptx")
        prs.save(safe_path)
        return safe_path
        
    def clean_metadata(self, file_path: str):
        prs = Presentation(file_path)
        prs.core_properties.author = "Ratchet Privacy Shield"
        prs.core_properties.last_modified_by = "Ratchet Privacy Shield"
        prs.core_properties.comments = ""
        prs.core_properties.title = ""
        prs.core_properties.subject = ""
        prs.save(file_path)
