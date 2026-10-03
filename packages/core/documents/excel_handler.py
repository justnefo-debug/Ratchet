import openpyxl
from .base_handler import DocumentHandler
from typing import List, Dict
from core.models import Entity
import os

class ExcelHandler(DocumentHandler):
    def extract_text(self, file_path: str) -> str:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        full_text = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value:
                        full_text.append(str(cell.value))
        wb.close()
        return "\n".join(full_text)
        
    def rebuild(self, file_path: str, mapping: Dict[str, str], entities: List[Entity]) -> str:
        wb = openpyxl.load_workbook(file_path)
        
        def replace_in_text(text: str) -> str:
            new_text = str(text)
            for placeholder, original in mapping.items():
                new_text = new_text.replace(original, placeholder)
            return new_text
            
        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        val_str = str(cell.value)
                        if any(orig in val_str for orig in mapping.values()):
                            cell.value = replace_in_text(val_str)
                            
        safe_path = file_path.replace(".xlsx", "_safe.xlsx")
        wb.save(safe_path)
        wb.close()
        return safe_path
        
    def clean_metadata(self, file_path: str):
        wb = openpyxl.load_workbook(file_path)
        wb.properties.creator = "Ratchet Privacy Shield"
        wb.properties.lastModifiedBy = "Ratchet Privacy Shield"
        wb.properties.title = ""
        wb.properties.subject = ""
        wb.save(file_path)
        wb.close()
