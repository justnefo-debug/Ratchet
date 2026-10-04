import docx
from .base_handler import DocumentHandler
from typing import List, Dict
from core.models import Entity
import os

class WordHandler(DocumentHandler):
    def extract_text(self, file_path: str) -> str:
        doc = docx.Document(file_path)
        full_text = []
        for para in doc.paragraphs:
            full_text.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    full_text.append(cell.text)
                    
        # Extract headers and footers
        for section in doc.sections:
            for header_para in section.header.paragraphs:
                full_text.append(header_para.text)
            for footer_para in section.footer.paragraphs:
                full_text.append(footer_para.text)
                
        return "\n".join(full_text)
        
    def rebuild(self, file_path: str, mapping: Dict[str, str], entities: List[Entity]) -> str:
        doc = docx.Document(file_path)
        
        def replace_in_text(text: str) -> str:
            new_text = text
            for placeholder, original in mapping.items():
                if original in new_text:
                    new_text = new_text.replace(original, placeholder)
            return new_text
            
        # Replace in paragraphs
        for para in doc.paragraphs:
            if any(orig in para.text for orig in mapping.values()):
                para.text = replace_in_text(para.text)
                
        # Replace in tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        if any(orig in para.text for orig in mapping.values()):
                            para.text = replace_in_text(para.text)
                            
        # Replace in headers and footers
        for section in doc.sections:
            for para in section.header.paragraphs:
                if any(orig in para.text for orig in mapping.values()):
                    para.text = replace_in_text(para.text)
            for para in section.footer.paragraphs:
                if any(orig in para.text for orig in mapping.values()):
                    para.text = replace_in_text(para.text)
                            
        safe_path = file_path.replace(".docx", "_safe.docx")
        doc.save(safe_path)
        return safe_path
        
    def clean_metadata(self, file_path: str):
        doc = docx.Document(file_path)
        prop = doc.core_properties
        prop.author = "Ratchet Privacy Shield"
        prop.last_modified_by = "Ratchet Privacy Shield"
        prop.comments = ""
        prop.title = ""
        prop.subject = ""
        doc.save(file_path)
