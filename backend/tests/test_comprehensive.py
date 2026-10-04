import sys
import os
import tempfile
import pandas as pd
from pptx import Presentation

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))) # backend
sys.path.append(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))) # root

from core.engine import PipelineEngine
from core.models import Entity
from packages.core.documents.text_handler import TextHandler
from packages.core.documents.csv_handler import CSVHandler
from packages.core.documents.pptx_handler import PPTXHandler

engine = PipelineEngine()

def test_api_logic():
    print("Testing Engine integration...")
    # Test Redact
    original = "My name is John Doe and my email is test@example.com."
    redacted, entities, sid = engine.process_redact(original)
    
    assert "[PERSON" in redacted
    assert "[EMAIL" in redacted
    assert sid is not None
    
    # Test Restore
    restored = engine.process_restore(redacted, sid)
    assert restored == original
    print("[OK] Engine process_redact and process_restore")

def test_document_handlers():
    print("Testing Document Handlers...")
    # 1. Text Handler
    with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode="w") as f:
        f.write("Contact alice@example.com.")
        temp_txt = f.name
        
    handler = TextHandler()
    text = handler.extract_text(temp_txt)
    assert "alice@example.com" in text
    
    redacted_text, entities, sid = engine.process_redact(text)
    mapping = engine.store.get_mapping(sid)
    safe_path = handler.rebuild(temp_txt, mapping, entities)
    
    with open(safe_path, "r") as f:
        safe_content = f.read()
    assert "alice@example.com" not in safe_content
    assert "[EMAIL" in safe_content
    os.remove(temp_txt)
    os.remove(safe_path)
    print("[OK] TextHandler")

    # 2. CSV Handler
    with tempfile.NamedTemporaryFile(suffix=".csv", delete=False, mode="w") as f:
        f.write("Name,Email\nBob,bob@test.com")
        temp_csv = f.name
        
    csv_h = CSVHandler()
    text = csv_h.extract_text(temp_csv)
    assert "bob@test.com" in text
    
    redacted_text, entities, sid = engine.process_redact(text)
    mapping = engine.store.get_mapping(sid)
    safe_csv = csv_h.rebuild(temp_csv, mapping, entities)
    
    df = pd.read_csv(safe_csv)
    assert df.iloc[0]["Email"] != "bob@test.com"
    os.remove(temp_csv)
    os.remove(safe_csv)
    print("[OK] CSVHandler")
    
    # 3. PPTX Handler
    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as f:
        temp_pptx = f.name
        
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Secret is secret@test.com"
    prs.save(temp_pptx)
    
    pptx_h = PPTXHandler()
    text = pptx_h.extract_text(temp_pptx)
    assert "secret@test.com" in text
    
    redacted_text, entities, sid = engine.process_redact(text)
    mapping = engine.store.get_mapping(sid)
    safe_pptx = pptx_h.rebuild(temp_pptx, mapping, entities)
    
    prs_safe = Presentation(safe_pptx)
    safe_text = prs_safe.slides[0].shapes.title.text
    assert "secret@test.com" not in safe_text
    os.remove(temp_pptx)
    os.remove(safe_pptx)
    print("[OK] PPTXHandler")

if __name__ == "__main__":
    test_api_logic()
    test_document_handlers()
    print("All backend logic checks passed!")
